import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock
from docx import Document
from transcription import Segment, discover, export, transcribe


class TranscriptionTests(unittest.TestCase):
    def test_discover_filters_media(self):
        with tempfile.TemporaryDirectory() as folder:
            for name in ["audio.M4V", "sound.m4a", "notes.md"]:
                Path(folder, name).touch()
            self.assertEqual([p.name for p in discover(folder)], ["audio.M4V", "sound.m4a"])

    def test_both_exports_unicode_timestamps_and_no_overwrite(self):
        with tempfile.TemporaryDirectory() as folder:
            segments = [Segment(65, 70, "Hola, ¿cómo estás?")]
            md, docx = export("audio.m4a", folder, segments, "es", True)
            self.assertIn(
                "[00:01:05 – 00:01:10] Hola, ¿cómo estás?", md.read_text(encoding="utf-8")
            )
            self.assertIn(
                "Hola, ¿cómo estás?", "\n".join(p.text for p in Document(docx).paragraphs)
            )
            second = export("audio.m4a", folder, segments, "en")
            self.assertNotEqual(md, second[0])
            self.assertTrue(md.exists() and docx.exists())

    def test_selected_output_formats(self):
        for formats in (("md",), ("docx",), ("md", "docx")):
            with self.subTest(formats=formats), tempfile.TemporaryDirectory() as folder:
                paths = export("audio.m4a", folder, [Segment(0, 1, "Hello")], "en", formats=formats)
                self.assertEqual({p.suffix for p in paths}, {"." + fmt for fmt in formats})
                # Windows temporary folders may use an 8.3 alias (RUNNER~1),
                # while export returns resolved, long-form paths.
                self.assertEqual(
                    {path.resolve() for path in Path(folder).iterdir()},
                    {path.resolve() for path in paths},
                )
                for path in paths:
                    text = (
                        path.read_text(encoding="utf-8")
                        if path.suffix == ".md"
                        else "\n".join(p.text for p in Document(path).paragraphs)
                    )
                    self.assertIn("Hello", text)

    def test_invalid_formats_write_nothing(self):
        with tempfile.TemporaryDirectory() as folder:
            for formats in ((), ("pdf",), ("md", "pdf")):
                with self.subTest(formats=formats), self.assertRaises(ValueError):
                    export("audio.m4a", folder, [], "en", formats=formats)
            self.assertEqual(list(Path(folder).iterdir()), [])

    def test_transcription_language_progress_and_silence(self):
        model = Mock()
        model.transcribe.return_value = (iter([Segment(0, 2, " Hello ")]), Mock(duration=2))
        progress = []
        result = transcribe(model, "audio.m4v", "en", progress.append)
        self.assertEqual(result[0].text, "Hello")
        self.assertEqual(model.transcribe.call_args.kwargs["language"], "en")
        self.assertEqual(progress[-1], 100)
        model.transcribe.return_value = (iter([]), Mock(duration=2))
        with self.assertRaisesRegex(ValueError, "No speech"):
            transcribe(model, "audio.m4v", "es")


if __name__ == "__main__":
    unittest.main()


class RobustnessTests(unittest.TestCase):
    def _run(self, segments, **kwargs):
        model = Mock()
        model.transcribe.return_value = (iter(segments), Mock(duration=10, language="es"))
        return model, transcribe(model, "a.wav", "es", **kwargs)

    def test_decoding_does_not_condition_on_previous_text(self):
        model, _ = self._run([Segment(0, 1, "Hola")])
        options = model.transcribe.call_args.kwargs
        self.assertIs(options["condition_on_previous_text"], False)
        self.assertEqual(options["no_speech_threshold"], 0.6)
        self.assertEqual(options["vad_parameters"]["min_silence_duration_ms"], 500)
        self.assertGreater(len(options["temperature"]), 1)
        self.assertIsNone(options["initial_prompt"])
        self.assertNotIn("hallucination_silence_threshold", options)

    def test_prompt_and_word_options_are_forwarded(self):
        model, _ = self._run(
            [Segment(0, 1, "SNS")], initial_prompt="  SNS, PDPCM ", word_timestamps=True
        )
        options = model.transcribe.call_args.kwargs
        self.assertEqual(options["initial_prompt"], "SNS, PDPCM")
        self.assertEqual(options["hallucination_silence_threshold"], 2.0)

    def test_confident_silence_is_dropped_but_quiet_speech_is_kept(self):
        from types import SimpleNamespace

        def seg(text, no_speech, logprob):
            return SimpleNamespace(
                start=0, end=1, text=text, words=None, no_speech_prob=no_speech, avg_logprob=logprob
            )

        _, result = self._run(
            [seg("Gracias por ver", 0.9, -1.5), seg("Hola", 0.9, -0.3), seg("Adiós", 0.1, -1.5)]
        )
        self.assertEqual([s.text for s in result], ["Hola", "Adiós"])
