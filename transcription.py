"""Local media transcription and paired document export."""

from dataclasses import dataclass
from pathlib import Path
from paths import MODEL_DIR
import os
import tempfile
import shutil

SUPPORTED = {
    ".m4v",
    ".m4a",
    ".mp4",
    ".mp3",
    ".wav",
    ".flac",
    ".ogg",
    ".aac",
    ".webm",
    ".mov",
    ".wma",
}


@dataclass
class Segment:
    start: float
    end: float
    text: str
    speaker: str | None = None
    words: list[tuple[float, float, str]] | None = None


def discover(folder: str | Path) -> list[Path]:
    folder = Path(folder).expanduser().resolve()
    if not folder.is_dir():
        raise ValueError("La carpeta de entrada no existe.")
    return sorted(
        (p for p in folder.iterdir() if p.is_file() and p.suffix.lower() in SUPPORTED),
        key=lambda p: p.name.casefold(),
    )


def timestamp(seconds: float) -> str:
    seconds = max(0, int(seconds))
    return f"{seconds // 3600:02}:{seconds // 60 % 60:02}:{seconds % 60:02}"


def srt_timestamp(seconds: float) -> str:
    milliseconds = max(0, round(seconds * 1000))
    whole, fraction = divmod(milliseconds, 1000)
    return f"{timestamp(whole)},{fraction:03d}"


PARAGRAPH_GAP = 1.5  # Seconds of silence that always start a new paragraph.
PARAGRAPH_SOFT_CHARS = 450  # After this length, break at the next sentence end.
PARAGRAPH_MAX_CHARS = 900  # Hard limit even without sentence punctuation.
SENTENCE_END = (".", "?", "!", "…", "。")


def group_paragraphs(segments: list[Segment]) -> list[Segment]:
    """Join short Whisper segments into readable paragraphs.

    A paragraph ends at a change of speaker, a long pause, or (once it is long
    enough) a sentence end. Of a run of identical segments, a typical decoding
    loop, only the first two are kept.
    """
    cleaned: list[Segment] = []
    run = 0
    for segment in segments:
        text = " ".join(segment.text.split())
        if not text:
            continue
        previous = cleaned[-1] if cleaned else None
        run = run + 1 if previous and previous.text.casefold() == text.casefold() else 1
        if run > 2:
            previous.end = max(previous.end, segment.end)
            continue
        cleaned.append(Segment(segment.start, segment.end, text, segment.speaker))
    paragraphs: list[Segment] = []
    for segment in cleaned:
        current = paragraphs[-1] if paragraphs else None
        if (
            current is not None
            and current.speaker == segment.speaker
            and segment.start - current.end < PARAGRAPH_GAP
            and len(current.text) + len(segment.text) < PARAGRAPH_MAX_CHARS
            and not (
                len(current.text) >= PARAGRAPH_SOFT_CHARS and current.text.endswith(SENTENCE_END)
            )
        ):
            current.text += " " + segment.text
            current.end = max(current.end, segment.end)
        else:
            paragraphs.append(segment)
    return paragraphs


def export(source, output, segments, language, timestamps=False, formats=("md", "docx")):
    formats = tuple(dict.fromkeys(formats))
    if not formats or any(fmt not in {"md", "docx", "srt"} for fmt in formats):
        raise ValueError("Select at least one valid output format: md, docx, srt.")
    # Keep output ordering stable for callers.
    extensions = tuple("." + fmt for fmt in ("md", "docx", "srt") if fmt in formats)
    output = Path(output).expanduser().resolve()
    output.mkdir(parents=True, exist_ok=True)
    # Include the source extension so similarly named audio/video files stay distinct.
    stem = Path(source).name
    title = Path(source).name
    language_name = {"es": "Español", "en": "English"}.get(language, language or "Auto")
    document = None
    if "docx" in formats:
        from docx import Document

        document = Document()
        document.add_heading(title, 0)
        document.add_paragraph(f"Idioma / Language: {language_name}")
    lines = [f"# {title}", "", f"Idioma / Language: {language_name}", ""]
    subtitles = []
    for segment in segments:
        text = segment.text.strip()
        if not text:
            continue
        if segment.speaker:
            text = f"{segment.speaker}: {text}"
        subtitles.append(
            f"{len(subtitles) + 1}\n{srt_timestamp(segment.start)} --> {srt_timestamp(segment.end)}\n{text}\n"
        )
    for paragraph in group_paragraphs(segments):
        text = paragraph.text
        if paragraph.speaker:
            text = f"{paragraph.speaker}: {text}"
        if timestamps:
            text = f"[{timestamp(paragraph.start)} – {timestamp(paragraph.end)}] {text}"
        lines.extend([text, ""])
        if document is not None:
            document.add_paragraph(text)
    temporary = []
    published = []
    try:
        for suffix in extensions:
            fd, name = tempfile.mkstemp(dir=output, suffix=suffix)
            os.close(fd)
            temporary.append(Path(name))
        for temp, suffix in zip(temporary, extensions):
            if suffix == ".md":
                temp.write_text("\n".join(lines), encoding="utf-8")
            elif suffix == ".srt":
                temp.write_text("\n".join(subtitles), encoding="utf-8")
            else:
                document.save(str(temp))
        index = 1
        while True:
            candidate = stem if index == 1 else f"{stem} ({index})"
            targets = tuple(output / f"{candidate}{ext}" for ext in extensions)
            reservation = output / f".{candidate}.export-lock"
            try:
                lock = reservation.open("xb")
            except FileExistsError:
                index += 1
                continue
            try:
                lock.close()
                if any(target.exists() for target in targets):
                    index += 1
                    continue
                try:
                    for temp, target in zip(temporary, targets):
                        # Atomic publication without replacing an existing file.
                        os.link(temp, target)
                        published.append(target)
                except FileExistsError:
                    for target in published:
                        target.unlink(missing_ok=True)
                    published.clear()
                    index += 1
                    continue
                except OSError as exc:
                    if exc.errno not in {1, 18, 38, 95}:
                        raise
                    # Filesystems without hard links retain exclusive-create safety.
                    for temp, target in zip(temporary, targets):
                        if target in published:
                            continue
                        with target.open("xb") as handle:
                            published.append(target)
                            with temp.open("rb") as source_handle:
                                shutil.copyfileobj(source_handle, handle)
                break
            finally:
                reservation.unlink(missing_ok=True)
    except Exception:
        for path in published:
            path.unlink(missing_ok=True)
        raise
    finally:
        for path in temporary:
            path.unlink(missing_ok=True)
    return targets


def load_model(name):
    from faster_whisper import WhisperModel

    return WhisperModel(name, device="cpu", compute_type="int8", download_root=str(MODEL_DIR))


def transcribe(
    model,
    source,
    language,
    on_progress=lambda value: None,
    word_timestamps=False,
    on_language=lambda language: None,
):
    segments, info = model.transcribe(
        str(source) if isinstance(source, (str, Path)) else source,
        language=language,
        task="transcribe",
        beam_size=5,
        vad_filter=True,
        word_timestamps=word_timestamps,
    )
    on_language(info.language)
    result = []
    for segment in segments:
        if segment.text.strip():
            words = (
                [(word.start, word.end, word.word) for word in segment.words]
                if word_timestamps and segment.words
                else None
            )
            result.append(Segment(segment.start, segment.end, segment.text.strip(), words=words))
        on_progress(min(100, segment.end / max(info.duration, 1) * 100))
    if not result:
        raise ValueError("No se ha detectado voz / No speech detected.")
    on_progress(100)
    return result
